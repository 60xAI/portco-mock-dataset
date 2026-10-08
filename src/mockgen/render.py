"""Renderer: content JSON -> real Office/PDF files. Deterministic; never calls a model."""
from __future__ import annotations

import copy
import io
import json
import math
import os
import random
import re
import shutil
import subprocess
import tempfile
import zipfile
import zlib
from datetime import datetime
from functools import lru_cache
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402

from .paths import ARCHIVE_ROOT, HELDBACK_ROOT, NUMBERS, CONTENT, TODAY  # noqa: E402

PH = re.compile(r"\{\{\s*(PRJ\d{4}):([^}]+?)\s*\}\}")


# ---------------------------------------------------------------- numbers access

@lru_cache(maxsize=512)
def numbers(pid: str) -> dict:
    p = NUMBERS / f"{pid}.json"
    if not p.exists():
        raise KeyError(f"no numbers for project {pid} (run `mockgen numbers`)")
    return json.loads(p.read_text())


def resolve_text(s, errors: list | None = None):
    if not isinstance(s, str) or "{{" not in s:
        return s

    def rep(m):
        pid, key = m.group(1), m.group(2).strip()
        try:
            v = numbers(pid)["values"].get(key)
        except KeyError as e:
            v = None
        if v is None:
            if errors is not None:
                errors.append(f"unknown placeholder {{{{{pid}:{key}}}}}")
            return m.group(0)
        return str(v)
    return PH.sub(rep, s)


def get_table(ref: str, errors: list | None = None) -> dict | None:
    try:
        pid, key = ref.split(":", 1)
        t = numbers(pid)["tables"].get(key)
    except (KeyError, ValueError):
        t = None
    if t is None and errors is not None:
        errors.append(f"unknown table_ref {ref}")
    return t


def table_data(tbl, errors=None) -> tuple[list[str], list[list], str | None]:
    """Returns (columns, rows, caption) for an inline table or a table_ref."""
    if tbl is None:
        return [], [], None
    if getattr(tbl, "table_ref", None):
        t = get_table(tbl.table_ref, errors)
        if not t:
            return [], [], tbl.caption
        cols, rows = list(t["columns"]), [list(r) for r in t["rows"]]
        if tbl.columns:
            cols = [cols[i] for i in tbl.columns if i < len(cols)]
            rows = [[r[i] for i in tbl.columns if i < len(r)] for r in rows]
        if tbl.max_rows:
            rows = rows[: tbl.max_rows]
        return cols, rows, tbl.caption or t.get("title")
    cols = [resolve_text(c, errors) for c in tbl.columns]
    rows = [[resolve_text(c, errors) if isinstance(c, str) else c for c in r] for r in tbl.rows]
    return cols, rows, tbl.caption


# ---------------------------------------------------------------- style

def hex2rgb(h: str):
    h = (h or "#333333").lstrip("#")
    if len(h) != 6:
        h = "333333"
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


class Style:
    def __init__(self, entry, world):
        self.entry = entry
        if entry.unit == "HO":
            t = world.group.template
            self.company = world.group.legal_name
        else:
            f = world.firms[entry.unit]
            temps = {x.id: x for x in f.templates}
            temps[world.group.template.id] = world.group.template
            t = temps.get(entry.template) or f.templates[-1]
            self.company = f.legacy_name if t.kind == "legacy" else (f.legacy_name if t.kind == "transition" else world.group.legal_name)
        self.t = t
        self.primary = t.primary_colour
        self.secondary = t.secondary_colour
        self.accent = t.accent_colour or t.secondary_colour
        self.font_h = t.font_heading
        self.font_b = t.font_body
        self.footer = t.footer_text
        self.brand = t.brand_line
        self.variant = zlib.crc32(t.id.encode()) % 3
        self.year = entry.created.year
        self.wide = self.year >= 2010


# ---------------------------------------------------------------- charts and molecules

def _num(x):
    if isinstance(x, (int, float)):
        return float(x)
    if isinstance(x, str):
        m = re.search(r"-?\d+(?:\.\d+)?", x.replace(",", ""))
        if m:
            return float(m.group(0))
    return None


def chart_png(ch, style: Style, errors=None, seed="x", size=(6.4, 3.8)) -> bytes:
    plt.rcParams.update({"font.family": "sans-serif", "font.size": 9})
    fig, ax = plt.subplots(figsize=size, dpi=130)
    cols = [np.array(hex2rgb(style.primary)) / 255, np.array(hex2rgb(style.secondary)) / 255, np.array(hex2rgb(style.accent)) / 255,
            (0.45, 0.45, 0.45), (0.8, 0.5, 0.1), (0.2, 0.6, 0.4)]
    rnd = random.Random(zlib.crc32(f"{seed}{ch.title}{ch.ref}".encode()))
    kind = ch.kind
    t = get_table(ch.ref, errors) if ch.ref else None
    if kind == "dose_response":
        conc = np.logspace(-2, 1.5, 7)
        if t and t["columns"] and "IC50" in " ".join(t["columns"]):
            mean_col = next((i for i, c in enumerate(t["columns"]) if "mean" in c.lower() or "IC50" in c), 1)
            pts = []
            for r in t["rows"]:
                v = _num(r[mean_col])
                if v is not None and not str(r[mean_col]).startswith(">"):
                    pts.append((r[0], v))
            pts = pts[:6] or [("cmpd", 3.0)]
        else:
            pts = [("Compound", 2.0)]
        xs = np.logspace(-2.2, 1.7, 100)
        for i, (name, ic) in enumerate(pts):
            ax.plot(xs, 100 / (1 + (ic / xs) ** -1 * 1) if False else 100 * xs / (xs + ic), color=cols[i % len(cols)], lw=1.4, label=f"{name} ({ic:g})")
            ys = 100 * conc / (conc + ic) + np.array([rnd.gauss(0, 3) for _ in conc])
            ax.scatter(conc, ys, color=cols[i % len(cols)], s=10)
        ax.set_xscale("log")
        ax.set_xlabel(ch.x_label or "Concentration (µM)")
        ax.set_ylabel(ch.y_label or "% inhibition")
        ax.set_ylim(-10, 110)
        ax.legend(fontsize=7, frameon=False)
    elif kind == "stability" and t:
        by = {}
        for r in t["rows"]:
            by.setdefault(r[0], []).append((_num(r[1]), _num(r[2])))
        for i, (cond, pts) in enumerate(by.items()):
            pts = [(x, y) for x, y in pts if x is not None and y is not None]
            ax.plot([p[0] for p in pts], [p[1] for p in pts], marker="o", color=cols[i % len(cols)], label=cond)
        ax.axhline(95, ls="--", color="grey", lw=0.8)
        ax.set_xlabel(ch.x_label or "Time (months)")
        ax.set_ylabel(ch.y_label or "Assay (% label claim)")
        ax.legend(fontsize=7, frameon=False)
    elif kind == "chromatogram":
        x = np.linspace(0, 20, 2000)
        y = np.zeros_like(x) + 0.5
        peaks = [(rnd.uniform(2, 18), rnd.uniform(5, 40), rnd.uniform(0.05, 0.12)) for _ in range(rnd.randint(3, 7))]
        peaks.append((rnd.uniform(8, 12), 400, 0.1))
        for c, h, wdt in peaks:
            y += h * np.exp(-((x - c) ** 2) / (2 * wdt ** 2))
        y += np.array([rnd.gauss(0, 0.3) for _ in x])
        ax.plot(x, y, color="black", lw=0.7)
        ax.set_xlabel("Time (min)")
        ax.set_ylabel("mAU")
    else:
        labels, series = ch.labels, ch.series
        if t is not None and ch.column is not None:
            rows = [r for r in t["rows"] if _num(r[ch.column]) is not None]
            labels = [str(r[0]) for r in rows][:30]
            series = [type("S", (), {"name": t["columns"][ch.column], "values": [_num(r[ch.column]) for r in rows][:30]})()]
        if not series:
            labels, series = ["n/a"], [type("S", (), {"name": "", "values": [0]})()]
        labels = labels or [str(i + 1) for i in range(len(series[0].values))]
        x = np.arange(len(labels))
        if kind in ("bar", "stacked_bar"):
            bottom = np.zeros(len(labels))
            wdt = 0.8 / (1 if kind == "stacked_bar" else len(series))
            for i, s in enumerate(series):
                vals = np.array((list(s.values) + [0] * len(labels))[: len(labels)], dtype=float)
                if kind == "stacked_bar":
                    ax.bar(x, vals, 0.6, bottom=bottom, color=cols[i % len(cols)], label=s.name)
                    bottom += vals
                else:
                    ax.bar(x + i * wdt - 0.4 + wdt / 2, vals, wdt, color=cols[i % len(cols)], label=s.name)
            ax.set_xticks(x, labels, rotation=45 if len(labels) > 6 else 0, ha="right" if len(labels) > 6 else "center", fontsize=7)
        elif kind == "hbar":
            s = series[0]
            ax.barh(x, (list(s.values) + [0] * len(labels))[: len(labels)], color=cols[0])
            ax.set_yticks(x, labels, fontsize=7)
        elif kind == "pie":
            s = series[0]
            ax.pie(s.values, labels=labels, colors=cols[: len(labels)], autopct="%1.0f%%", textprops={"fontsize": 7})
        elif kind == "scatter":
            s = series[0]
            ys = series[1].values if len(series) > 1 else s.values
            xs = s.values if len(series) > 1 else list(range(len(s.values)))
            ax.scatter(xs, ys, color=cols[0], s=14)
        else:
            for i, s in enumerate(series):
                ax.plot(labels, (list(s.values) + [None] * len(labels))[: len(labels)], marker="o", color=cols[i % len(cols)], label=s.name)
        if len(series) > 1 and kind != "pie":
            ax.legend(fontsize=7, frameon=False)
        if ch.x_label:
            ax.set_xlabel(ch.x_label)
        if ch.y_label:
            ax.set_ylabel(ch.y_label)
    if ch.title:
        ax.set_title(ch.title, fontsize=10)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def molecule_png(ref: str, n: int, errors=None) -> bytes:
    from rdkit import Chem
    from rdkit.Chem import Draw
    smis, legends = [], []
    if ref and ref.startswith("PRJ"):
        try:
            cs = numbers(ref.split(":")[0])["compounds"]
        except KeyError:
            cs = []
            if errors is not None:
                errors.append(f"molecule_grid: unknown project {ref}")
        for c in cs[:n]:
            smis.append(c["smiles"]); legends.append(c["id"])
    elif ref and ref.startswith("SER"):
        allc = json.loads((NUMBERS / "compounds.json").read_text())
        for c in allc.get(ref, [])[:n]:
            smis.append(c["smiles"]); legends.append(c["id"])
    if not smis:
        if errors is not None:
            errors.append(f"molecule_grid: no compounds for {ref}")
        return chart_png_placeholder()
    mols = [Chem.MolFromSmiles(s) for s in smis]
    img = Draw.MolsToGridImage(mols, molsPerRow=min(4, len(mols)), subImgSize=(260, 200), legends=legends)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def chart_png_placeholder() -> bytes:
    fig, ax = plt.subplots(figsize=(3, 2))
    ax.axis("off")
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def boxes_png(kind: str, boxes: list[str], style: Style) -> bytes:
    fig, ax = plt.subplots(figsize=(7, 3.6), dpi=130)
    ax.axis("off")
    col = np.array(hex2rgb(style.primary)) / 255
    if kind == "timeline":
        n = len(boxes)
        ax.plot([0, 1], [0.5, 0.5], color="grey", lw=2)
        for i, b in enumerate(boxes):
            x = (i + 0.5) / n
            ax.scatter([x], [0.5], s=80, color=col)
            ax.text(x, 0.62 if i % 2 == 0 else 0.3, b, ha="center", va="center", fontsize=7, wrap=True)
    else:
        parents = {}
        for b in boxes:
            if ">" in b:
                p, c = [x.strip() for x in b.split(">", 1)]
                parents.setdefault(p, []).append(c)
            else:
                parents.setdefault(b.strip(), [])
        roots = [p for p in parents if not any(p in v for v in parents.values())]
        levels = []
        cur = roots
        seen = set()
        while cur:
            levels.append(cur)
            seen.update(cur)
            nxt = []
            for p in cur:
                nxt += [c for c in parents.get(p, []) if c not in seen]
            cur = nxt
        pos = {}
        for li, lvl in enumerate(levels):
            for i, name in enumerate(lvl):
                x = (i + 0.5) / len(lvl)
                y = 1 - (li + 0.5) / len(levels)
                pos[name] = (x, y)
                ax.text(x, y, name, ha="center", va="center", fontsize=6.5, color="white",
                        bbox=dict(boxstyle="round,pad=0.4", fc=col, ec="none"))
        for p, cs in parents.items():
            for c in cs:
                if p in pos and c in pos:
                    ax.plot([pos[p][0], pos[c][0]], [pos[p][1] - 0.04, pos[c][1] + 0.04], color="grey", lw=0.6, zorder=0)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def image_png(img, style, errors=None, seed="x") -> bytes:
    if img.kind == "molecule_grid":
        return molecule_png(img.ref, img.n, errors)
    if img.kind == "chart" and img.chart:
        return chart_png(img.chart, style, errors, seed)
    if img.kind in ("org_chart", "timeline"):
        return boxes_png(img.kind, img.boxes or [], style)
    fig, ax = plt.subplots(figsize=(4, 1))
    ax.axis("off")
    ax.text(0.5, 0.5, img.text or style.brand, ha="center", va="center", fontsize=16, color=np.array(hex2rgb(style.primary)) / 255, fontweight="bold")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------- properties

def app_meta(fmt: str, year: int):
    if year < 2001:
        ver = ("9.0", "2000")
    elif year < 2003:
        ver = ("10.0", "XP")
    elif year < 2007:
        ver = ("11.0", "2003")
    elif year < 2010:
        ver = ("12.0", "2007")
    elif year < 2013:
        ver = ("14.0", "2010")
    elif year < 2016:
        ver = ("15.0", "2013")
    else:
        ver = ("16.0", "2016")
    app = {"pptx": "Microsoft Office PowerPoint", "docx": "Microsoft Office Word", "xlsx": "Microsoft Excel"}[fmt]
    if year >= 2010 and fmt != "xlsx":
        app = app.replace("Office ", "")
    return app, ver


def patch_app_xml(path: Path, company: str, fmt: str, year: int):
    app, (ver, _) = app_meta(fmt, year)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                s = data.decode("utf-8")
                s = re.sub(r"<Company>.*?</Company>|<Company/>", "", s)
                s = re.sub(r"<Application>.*?</Application>", f"<Application>{app}</Application>", s)
                if "<Application>" not in s:
                    s = s.replace("</Properties>", f"<Application>{app}</Application></Properties>")
                s = re.sub(r"<AppVersion>.*?</AppVersion>", "", s)
                s = s.replace("</Properties>", f"<Company>{_xml(company)}</Company><AppVersion>{ver}000</AppVersion></Properties>")
                data = s.encode("utf-8")
            zout.writestr(item, data)
        if "docProps/app.xml" not in names:
            pass
    tmp.replace(path)


def _xml(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def set_core(cp, entry, people, props, title_default):
    cp.author = people(entry.author)
    cp.last_modified_by = people(entry.last_saved_by)
    cp.created = entry.created
    cp.modified = entry.modified
    cp.title = (props.title if props and props.title else title_default) or ""
    if props:
        if props.subject:
            cp.subject = props.subject
        if props.keywords:
            cp.keywords = props.keywords
        if props.category:
            cp.category = props.category
        if props.comments:
            cp.comments = props.comments
    try:
        cp.revision = 1 + (zlib.crc32(entry.id.encode()) % 17)
    except Exception:
        pass


def set_pdf_meta(path: Path, entry, people, kind: str, title: str):
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(str(path))
    w = PdfWriter()
    for pg in r.pages:
        w.add_page(pg)
    y = entry.created.year
    _, (_, label) = app_meta("docx", y)
    if kind == "scan":
        creator, producer = "Scan to PDF", "MFP PDF Library 2.1"
        author = ""
    elif kind == "export":
        creator, producer = f"Microsoft® PowerPoint® {label}" if y >= 2007 else "PScript5.dll Version 5.2.2", f"Microsoft® PowerPoint® {label}" if y >= 2007 else "Acrobat Distiller 6.0 (Windows)"
        author = people(entry.author)
    else:
        creator, producer = (f"Microsoft® Word {label}", f"Microsoft® Word {label}") if y >= 2007 else ("PScript5.dll Version 5.2.2", "Acrobat Distiller 7.0.5 (Windows)")
        author = people(entry.author)

    def pdfdate(d: datetime):
        return d.strftime("D:%Y%m%d%H%M%S+00'00'")
    meta = {"/Creator": creator, "/Producer": producer, "/CreationDate": pdfdate(entry.created), "/ModDate": pdfdate(entry.modified)}
    if author:
        meta["/Author"] = author
    if title and kind != "scan":
        meta["/Title"] = title
    w.add_metadata(meta)
    with open(path, "wb") as fh:
        w.write(fh)


def set_mtime(path: Path, entry):
    ts = entry.modified.timestamp()
    os.utime(path, (ts, ts))


# ---------------------------------------------------------------- LibreOffice

def lo_convert(src: Path, fmt: str, outdir: Path, timeout=240) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    prof = tempfile.mkdtemp(prefix="lo_prof_")
    try:
        cmd = ["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--norestore", "--convert-to", fmt, "--outdir", str(outdir), str(src)]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        out = outdir / (src.stem + "." + fmt.split(":")[0])
        if not out.exists():
            raise RuntimeError(f"LibreOffice failed converting {src.name} to {fmt}: {res.stderr.strip()[:400] or res.stdout.strip()[:400]}")
        return out
    finally:
        shutil.rmtree(prof, ignore_errors=True)


# ---------------------------------------------------------------- PPTX

def render_deck(content, entry, style: Style, people, out: Path, errors: list):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt, Emu

    prs = Presentation()
    if style.wide:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    prim, sec = RGBColor(*hex2rgb(style.primary)), RGBColor(*hex2rgb(style.secondary))
    L = {"title": 0, "title_content": 1, "section": 2, "two_content": 3, "title_only": 5, "blank": 6, "table": 5, "chart": 5, "image": 5, "quote": 5}

    def font(run, size=None, bold=None, color=None, heading=False):
        run.font.name = style.font_h if heading else style.font_b
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color

    def bullets_into(tf, items, base=20):
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        n = len(items)
        base = base if n <= 6 else (base - 2 if n <= 9 else base - 4)
        first = True
        for b in items:
            text = b if isinstance(b, str) else b.text
            lvl = 0 if isinstance(b, str) else b.level
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.level = min(lvl, 4)
            r = para.add_run()
            r.text = resolve_text(text, errors)
            font(r, size=max(12, base - 3 * lvl))

    for idx, sd in enumerate(content.slides, 1):
        layout = prs.slide_layouts[L[sd.layout]]
        s = prs.slides.add_slide(layout)
        title_ph = s.shapes.title
        # reposition placeholders for wide slides
        for ph in s.placeholders:
            if style.wide:
                ph.left = int(ph.left * 13.333 / 10)
                ph.width = int(ph.width * 13.333 / 10)
        # title band / rule
        if sd.layout not in ("title", "section", "blank"):
            if style.variant == 0:
                band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.15))
                band.fill.solid(); band.fill.fore_color.rgb = prim; band.line.fill.background()
                s.shapes._spTree.remove(band._element); s.shapes._spTree.insert(2, band._element)
                tcolor = RGBColor(255, 255, 255)
            elif style.variant == 1:
                bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.2), W - Inches(0.8), Emu(38100))
                bar.fill.solid(); bar.fill.fore_color.rgb = sec; bar.line.fill.background()
                tcolor = prim
            else:
                side = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), H)
                side.fill.solid(); side.fill.fore_color.rgb = prim; side.line.fill.background()
                tcolor = prim
        elif sd.layout in ("title", "section"):
            bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
            bg.fill.solid(); bg.fill.fore_color.rgb = prim if sd.layout == "title" or style.variant != 1 else sec
            bg.line.fill.background()
            s.shapes._spTree.remove(bg._element); s.shapes._spTree.insert(2, bg._element)
            tcolor = RGBColor(255, 255, 255)
        else:
            tcolor = prim
        if sd.layout in ("title", "section"):
            for ph in s.placeholders:
                idx = ph.placeholder_format.idx
                ph.left, ph.width = Inches(0.8), W - Inches(1.6)
                if idx == 0:
                    ph.top, ph.height = int(H * 0.30), Inches(1.6)
                elif idx == 1:
                    ph.top, ph.height = int(H * 0.30) + Inches(1.7), Inches(1.2)
        if title_ph is not None:
            if sd.layout not in ("title", "section"):
                title_ph.top, title_ph.height = Inches(0.15), Inches(0.9)
                title_ph.left, title_ph.width = Inches(0.45), W - Inches(0.9)
            tf = title_ph.text_frame
            tf.clear()
            r = tf.paragraphs[0].add_run()
            r.text = resolve_text(sd.title or "", errors)
            font(r, size=40 if sd.layout == "title" else (32 if sd.layout == "section" else 26), bold=True, color=tcolor, heading=True)
            if sd.layout not in ("title", "section"):
                tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        elif sd.title:
            tb = s.shapes.add_textbox(Inches(0.45), Inches(0.2), W - Inches(0.9), Inches(0.9))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = resolve_text(sd.title, errors)
            font(r, size=26, bold=True, color=tcolor, heading=True)
        body_top = Inches(1.4)
        body_h = H - Inches(2.1)
        if sd.layout == "title":
            sub = [p for p in s.placeholders if p.placeholder_format.idx == 1]
            if sub:
                tf = sub[0].text_frame; tf.clear()
                r = tf.paragraphs[0].add_run(); r.text = resolve_text(sd.subtitle or style.brand, errors)
                font(r, size=20, color=RGBColor(235, 235, 235))
        elif sd.layout == "section":
            for p in s.placeholders:
                if p.placeholder_format.idx == 1:
                    tf = p.text_frame; tf.clear()
                    r = tf.paragraphs[0].add_run(); r.text = resolve_text(sd.subtitle or "", errors)
                    font(r, size=18, color=RGBColor(230, 230, 230))
        elif sd.layout == "title_content":
            body = [p for p in s.placeholders if p.placeholder_format.idx == 1][0]
            body.top, body.height = body_top, body_h
            body.left, body.width = Inches(0.5), W - Inches(1.0)
            if sd.bullets:
                bullets_into(body.text_frame, sd.bullets)
            else:
                body._element.getparent().remove(body._element)
        elif sd.layout == "two_content":
            phs = [p for p in s.placeholders if p.placeholder_format.idx in (1, 2)]
            for ph, items, left in zip(phs, [sd.bullets, sd.bullets_right], [Inches(0.5), W // 2 + Inches(0.1)]):
                ph.top, ph.height, ph.left, ph.width = body_top, body_h, left, W // 2 - Inches(0.6)
                if items:
                    bullets_into(ph.text_frame, items, base=18)
                else:
                    ph._element.getparent().remove(ph._element)
        elif sd.layout == "quote":
            tb = s.shapes.add_textbox(Inches(1), Inches(2), W - Inches(2), Inches(3))
            tb.text_frame.word_wrap = True
            r = tb.text_frame.paragraphs[0].add_run(); r.text = resolve_text(" ".join(b if isinstance(b, str) else b.text for b in sd.bullets), errors)
            font(r, size=24, color=prim); r.font.italic = True
        # free-form content for title_only-based layouts
        y = body_top
        avail_w = W - Inches(1.0)
        if sd.layout in ("table", "chart", "image", "title_only") and sd.bullets:
            tb = s.shapes.add_textbox(Inches(0.5), y, avail_w, Inches(1.2))
            tb.text_frame.word_wrap = True
            bullets_into(tb.text_frame, sd.bullets, base=16)
            y += Inches(min(2.0, 0.35 * len(sd.bullets) + 0.2))
        if sd.table is not None:
            cols, rows, cap = table_data(sd.table, errors)
            if cols:
                rows = rows[:16]
                nrows = len(rows) + 1
                th = min(H - y - Inches(0.7), Inches(0.3) * nrows)
                shape = s.shapes.add_table(nrows, len(cols), Inches(0.5), y, avail_w, th)
                tbl = shape.table
                fs = 14 if nrows <= 8 else (12 if nrows <= 13 else 10)
                if len(cols) > 7:
                    fs -= 1
                for j, c in enumerate(cols):
                    cell = tbl.cell(0, j)
                    cell.text = str(c)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            font(r, size=fs, bold=True, color=RGBColor(255, 255, 255))
                    cell.fill.solid(); cell.fill.fore_color.rgb = prim
                for i, row in enumerate(rows, 1):
                    for j in range(len(cols)):
                        cell = tbl.cell(i, j)
                        cell.text = "" if j >= len(row) or row[j] is None else str(resolve_text(row[j], errors) if isinstance(row[j], str) else row[j])
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                font(r, size=fs)
                y += th + Inches(0.1)
        pic_png = None
        if sd.chart is not None:
            pic_png = chart_png(sd.chart, style, errors, seed=f"{entry.id}-{idx}")
        elif sd.image is not None:
            pic_png = image_png(sd.image, style, errors, seed=f"{entry.id}-{idx}")
        if pic_png:
            from PIL import Image as PImage
            im = PImage.open(io.BytesIO(pic_png))
            maxh = H - y - Inches(0.6)
            maxw = avail_w if sd.layout != "title_content" or not sd.bullets else W // 2 - Inches(0.5)
            left = Inches(0.5) if not (sd.layout == "title_content" and sd.bullets) else W // 2
            ratio = im.width / im.height
            wdt = min(maxw, int(maxh * ratio))
            s.shapes.add_picture(io.BytesIO(pic_png), int(left + (maxw - wdt) // 2), int(y), width=int(wdt))
            if sd.layout == "title_content" and sd.bullets:
                body = [p for p in s.placeholders if p.placeholder_format.idx == 1]
                if body:
                    body[0].width = W // 2 - Inches(0.6)
        # footer
        if sd.layout not in ("title",):
            ft = s.shapes.add_textbox(Inches(0.4), H - Inches(0.45), W - Inches(0.8), Inches(0.3))
            p = ft.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = resolve_text(sd.footer or f"{style.footer}", errors) + f"    {idx}"
            font(r, size=8, color=RGBColor(110, 110, 110))
        if sd.notes:
            s.notes_slide.notes_text_frame.text = resolve_text(sd.notes, errors)
        if sd.hidden:
            s._element.set("show", "0")
    set_core(prs.core_properties, entry, people, content.properties, content.slides[0].title if content.slides else entry.filename)
    prs.save(out)
    patch_app_xml(out, style.company, "pptx", entry.created.year)


# ---------------------------------------------------------------- XLSX

def render_workbook(content, entry, style: Style, people, out: Path, errors: list):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.utils.cell import coordinate_from_string, range_boundaries

    wb = Workbook()
    wb.remove(wb.active)
    hfill = PatternFill("solid", fgColor=style.primary.lstrip("#"))
    fname = style.font_b if style.year >= 2007 else "Arial"
    for sh in content.sheets:
        name = re.sub(r"[\[\]\*\?/\\:]", "-", sh.name)[:31] or "Sheet1"
        ws = wb.create_sheet(name)
        col0, row0 = coordinate_from_string(sh.start)
        c0 = column_index_from_string(col0)
        for i, row in enumerate(sh.rows):
            for j, v in enumerate(row):
                cell = ws.cell(row=row0 + i, column=c0 + j)
                if isinstance(v, dict):
                    from .schemas import CellObj
                    try:
                        co = CellObj.model_validate(v)
                    except Exception as e:
                        errors.append(f"sheet {sh.name} row {i + 1} col {j + 1}: bad cell object: {e}")
                        continue
                    val = co.v
                    if co.ref:
                        pid, key = co.ref.split(":", 1) if ":" in co.ref else (co.ref, "")
                        try:
                            wv = numbers(pid)["values"].get(key)
                        except KeyError:
                            wv = None
                        if wv is None:
                            errors.append(f"sheet {sh.name}: unknown ref {co.ref}")
                        elif val is None:
                            val = _coerce(wv)
                        elif str(val).replace(",", "") != str(wv).replace(",", "") and _num(val) != _num(wv):
                            errors.append(f"sheet {sh.name}: cell value {val!r} contradicts world value {wv!r} for {co.ref}")
                    cell.value = resolve_text(val, errors) if isinstance(val, str) else val
                    cell.font = Font(name=fname, bold=co.bold, italic=co.italic, color=(co.color or "000000").lstrip("#"))
                    if co.fill:
                        cell.fill = PatternFill("solid", fgColor=co.fill.lstrip("#"))
                    if co.fmt:
                        cell.number_format = co.fmt
                else:
                    cell.value = resolve_text(v, errors) if isinstance(v, str) else v
                    cell.font = Font(name=fname, bold=(i == 0 and isinstance(v, str) and len(sh.rows) > 1))
        for dr in sh.data_refs:
            t = get_table(dr.table_ref, errors)
            if not t:
                continue
            cols, rows = list(t["columns"]), [list(r) for r in t["rows"]]
            if dr.columns:
                cols = [cols[i] for i in dr.columns if i < len(cols)]
                rows = [[r[i] for i in dr.columns if i < len(r)] for r in rows]
            if dr.max_rows:
                rows = rows[: dr.max_rows]
            cl, rw = coordinate_from_string(dr.at)
            cc = column_index_from_string(cl)
            r0 = rw
            if dr.header:
                for j, c in enumerate(cols):
                    cell = ws.cell(row=rw, column=cc + j, value=c)
                    cell.font = Font(name=fname, bold=True, color="FFFFFF")
                    cell.fill = hfill
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                r0 = rw + 1
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    ws.cell(row=r0 + i, column=cc + j, value=_coerce(v)).font = Font(name=fname)
            for note_i, note in enumerate(t.get("notes", [])):
                ws.cell(row=r0 + len(rows) + 1 + note_i, column=cc, value=note).font = Font(name=fname, italic=True, size=9)
            for j, c in enumerate(cols):
                L_ = get_column_letter(cc + j)
                if L_ not in sh.col_widths:
                    ws.column_dimensions[L_].width = max(ws.column_dimensions[L_].width or 8, min(28, max(10, len(str(c)) * 0.9)))
        for m in sh.merged:
            try:
                ws.merge_cells(m)
            except Exception as e:
                errors.append(f"sheet {sh.name}: bad merged range {m}: {e}")
        for col, wd in sh.col_widths.items():
            ws.column_dimensions[col].width = wd
        for rng, fmt in sh.number_formats.items():
            try:
                mnc, mnr, mxc, mxr = range_boundaries(rng)
                for r in ws.iter_rows(min_row=mnr, max_row=mxr, min_col=mnc, max_col=mxc):
                    for c in r:
                        c.number_format = fmt
            except Exception as e:
                errors.append(f"sheet {sh.name}: bad number_format range {rng}: {e}")
        for cm in sh.comments:
            try:
                ws[cm.cell].comment = Comment(resolve_text(cm.text, errors), cm.author or people(entry.author))
            except Exception as e:
                errors.append(f"sheet {sh.name}: bad comment cell {cm.cell}: {e}")
        if sh.freeze:
            ws.freeze_panes = sh.freeze
        if sh.hidden:
            ws.sheet_state = "hidden"
        if sh.tab_color:
            ws.sheet_properties.tabColor = sh.tab_color.lstrip("#")
        if sh.chart:
            png = chart_png(sh.chart, style, errors, seed=f"{entry.id}-{sh.name}", size=(6, 3.4))
            img = XLImage(io.BytesIO(png))
            img.width, img.height = 600, 340
            ws.add_image(img, sh.chart_at or "J2")
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    if all(wb[s].sheet_state == "hidden" for s in wb.sheetnames):
        wb[wb.sheetnames[0]].sheet_state = "visible"
    p = wb.properties
    p.creator = people(entry.author)
    p.lastModifiedBy = people(entry.last_saved_by)
    p.created = entry.created
    p.modified = entry.modified
    p.title = content.properties.title or ""
    if content.properties.subject:
        p.subject = content.properties.subject
    if content.properties.keywords:
        p.keywords = content.properties.keywords
    wb.save(out)
    patch_app_xml(out, style.company, "xlsx", entry.created.year)


def _coerce(v):
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return v


# ---------------------------------------------------------------- DOCX

def render_document(content, entry, style: Style, people, out: Path, errors: list):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches, Cm

    doc = Document()
    sec = doc.sections[0]
    if style.entry.unit != "HO" and style_is_us(style):
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    else:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if content.orientation == "landscape":
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(2.2)
    prim = RGBColor(*hex2rgb(style.primary))
    st = doc.styles["Normal"]
    st.font.name = style.font_b
    st.font.size = Pt(10.5 if style.font_b.lower() not in ("times new roman", "georgia", "garamond") else 11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), style.font_b)
    for lvl in (1, 2, 3):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name = style.font_h
        hs.font.color.rgb = prim
        hs.font.size = Pt({1: 15, 2: 12.5, 3: 11}[lvl])
    ts = doc.styles["Title"]
    ts.font.name = style.font_h
    ts.font.color.rgb = prim
    ts.font.size = Pt(24)

    if content.header:
        hp = sec.header.paragraphs[0]
        hp.text = resolve_text(content.header, errors)
        for r in hp.runs:
            r.font.size = Pt(8); r.font.color.rgb = RGBColor(100, 100, 100)
    fp = sec.footer.paragraphs[0]
    fp.text = resolve_text(content.footer or style.footer, errors)
    for r in fp.runs:
        r.font.size = Pt(8); r.font.color.rgb = RGBColor(100, 100, 100)
    _add_page_number(fp)
    if content.watermark:
        _add_watermark(sec, content.watermark)

    if content.title_block:
        tb = content.title_block
        brand = doc.add_paragraph()
        r = brand.add_run(style.brand); r.font.size = Pt(9); r.font.color.rgb = prim; r.bold = True
        doc.add_paragraph(resolve_text(tb.title, errors), style="Title")
        if tb.subtitle:
            sp = doc.add_paragraph()
            r = sp.add_run(resolve_text(tb.subtitle, errors)); r.font.size = Pt(13); r.italic = True
        for ln in tb.lines:
            doc.add_paragraph(resolve_text(ln, errors))

    def add_blocks(blocks):
        for b in blocks:
            if b.page_break:
                doc.add_page_break()
            if b.p is not None:
                para = doc.add_paragraph()
                r = para.add_run(resolve_text(b.p, errors))
                if b.style == "bold":
                    r.bold = True
                elif b.style == "quote":
                    r.italic = True; para.paragraph_format.left_indent = Cm(1)
                elif b.style == "note":
                    r.italic = True; r.font.size = Pt(9)
                elif b.style == "small":
                    r.font.size = Pt(8.5)
            if b.items:
                for it in b.items:
                    text = it if isinstance(it, str) else it.text
                    lvl = 0 if isinstance(it, str) else it.level
                    sname = ("List Number" if b.ordered else "List Bullet") + (f" {lvl + 1}" if lvl else "")
                    try:
                        doc.add_paragraph(resolve_text(text, errors), style=sname)
                    except KeyError:
                        doc.add_paragraph(resolve_text(text, errors), style="List Bullet")
            if b.table is not None:
                cols, rows, cap = table_data(b.table, errors)
                if cols:
                    if cap:
                        cp = doc.add_paragraph()
                        r = cp.add_run(resolve_text(cap, errors)); r.bold = True; r.font.size = Pt(9)
                    t = doc.add_table(rows=1 + len(rows), cols=len(cols))
                    t.style = "Table Grid"
                    fs = Pt(8 if len(cols) > 6 else 9)
                    for j, c in enumerate(cols):
                        cell = t.cell(0, j)
                        cell.text = str(c)
                        _shade(cell, style.primary)
                        for p_ in cell.paragraphs:
                            for r in p_.runs:
                                r.bold = True; r.font.size = fs; r.font.color.rgb = RGBColor(255, 255, 255)
                    for i, row in enumerate(rows, 1):
                        for j in range(len(cols)):
                            v = row[j] if j < len(row) else ""
                            cell = t.cell(i, j)
                            cell.text = "" if v is None else str(resolve_text(v, errors) if isinstance(v, str) else v)
                            for p_ in cell.paragraphs:
                                for r in p_.runs:
                                    r.font.size = fs
                    if getattr(b.table, "table_ref", None):
                        tt = get_table(b.table.table_ref)
                        for note in (tt or {}).get("notes", []):
                            np_ = doc.add_paragraph()
                            r = np_.add_run(note); r.italic = True; r.font.size = Pt(8)
                    doc.add_paragraph()
            if b.chart is not None:
                doc.add_picture(io.BytesIO(chart_png(b.chart, style, errors, seed=entry.id)), width=Inches(5.8))
            if b.image is not None:
                doc.add_picture(io.BytesIO(image_png(b.image, style, errors, seed=entry.id)), width=Inches(5.6))

    for s in content.sections:
        if s.heading:
            doc.add_heading(resolve_text(s.heading, errors), level=max(1, min(3, s.level)))
        add_blocks(s.blocks)
    if content.signature_block:
        doc.add_paragraph()
        t = doc.add_table(rows=1 + len(content.signature_block), cols=4)
        t.style = "Table Grid"
        for j, h in enumerate(["Name", "Role", "Signature", "Date"]):
            t.cell(0, j).text = h
        for i, sg in enumerate(content.signature_block, 1):
            t.cell(i, 0).text = resolve_text(sg.name, errors)
            t.cell(i, 1).text = resolve_text(sg.role, errors)
            if sg.signed:
                p_ = t.cell(i, 2).paragraphs[0]
                r = p_.add_run(sg.name.split()[-1] if sg.name else ""); r.italic = True; r.font.name = "Georgia"
            t.cell(i, 3).text = sg.date or ""
    for a in content.appendices:
        doc.add_page_break()
        if a.heading:
            doc.add_heading(resolve_text(a.heading, errors), level=1)
        add_blocks(a.blocks)
    for tc in content.tracked_changes:
        if not _tracked_change(doc, tc):
            errors.append(f"tracked_changes: text not found in a single paragraph: {tc.find[:60]!r}")
    for cm in content.comments:
        ok = False
        for para in _all_paragraphs(doc):
            if cm.anchor and cm.anchor in para.text and para.runs:
                try:
                    doc.add_comment(para.runs, text=resolve_text(cm.text, errors), author=cm.author, initials="".join(w[0] for w in cm.author.split()[:2]))
                    ok = True
                except Exception as e:
                    errors.append(f"comment failed: {e}")
                    ok = True
                break
        if not ok:
            errors.append(f"comments: anchor text not found in a single paragraph: {cm.anchor[:60]!r}")
    title = content.title_block.title if content.title_block else (content.properties.title or entry.filename)
    set_core(doc.core_properties, entry, people, content.properties, resolve_text(title))
    doc.save(out)
    patch_app_xml(out, style.company, "docx", entry.created.year)


def style_is_us(style: Style) -> bool:
    from .world import world
    try:
        return world().firms[style.entry.unit].date_convention == "US"
    except Exception:
        return False


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def _shade(cell, hexcol):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcol.lstrip("#"))
    tcPr.append(shd)


def _add_page_number(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run("    Page ")
    r2 = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        r2._r.append(el)


def _add_watermark(section, text):
    from docx.oxml import parse_xml
    hdr = section.header
    p = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
        '<w:pict><v:shape id="PowerPlusWaterMarkObject" o:spid="_x0000_s2049" type="#_x0000_t136" '
        'style="position:absolute;margin-left:0;margin-top:0;width:468pt;height:117pt;rotation:315;z-index:-251654144;'
        'mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" '
        'o:allowincell="f" fillcolor="silver" stroked="f"><v:fill opacity=".5"/>'
        f'<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{_xml(text)}"/></v:shape></w:pict></w:r>'
    )
    p._p.append(parse_xml(xml))


def _tracked_change(doc, tc) -> bool:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for para in _all_paragraphs(doc):
        if tc.find not in para.text:
            continue
        full = para.text
        i = full.index(tc.find)
        before, target, after = full[:i], full[i:i + len(tc.find)], full[i + len(tc.find):]
        rpr = copy.deepcopy(para.runs[0]._r.rPr) if para.runs and para.runs[0]._r.rPr is not None else None
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        date = (tc.date or TODAY.isoformat()) + "T09:00:00Z"

        def run(text, deleted=False):
            r = OxmlElement("w:r")
            if rpr is not None:
                r.append(copy.deepcopy(rpr))
            t = OxmlElement("w:delText" if deleted else "w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r.append(t)
            return r

        def wrap(tag, r, n):
            w = OxmlElement(tag)
            w.set(qn("w:id"), str(n)); w.set(qn("w:author"), tc.author); w.set(qn("w:date"), date)
            w.append(r)
            return w
        n = zlib.crc32(tc.find.encode()) % 100000
        p = para._p
        if before:
            p.append(run(before))
        if tc.delete:
            p.append(wrap("w:del", run(target, True), n))
        else:
            p.append(run(target))
        if tc.insert:
            p.append(wrap("w:ins", run(tc.insert), n + 1))
        if after:
            p.append(run(after))
        return True
    return False


# ---------------------------------------------------------------- scan

def make_scan(pdf_in: Path, pdf_out: Path, settings, seed: str):
    from pdf2image import convert_from_path
    from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps
    rnd = random.Random(zlib.crc32(seed.encode()))
    pages = convert_from_path(str(pdf_in), dpi=200, grayscale=True)
    if settings.pages:
        pages = pages[: settings.pages]
    out = []
    for i, im in enumerate(pages):
        im = im.convert("L")
        skew = settings.skew_deg + rnd.uniform(-0.25, 0.25)
        skew = max(-1.5, min(1.5, skew))
        im = im.rotate(skew, resample=Image.BICUBIC, expand=False, fillcolor=255)
        draw = ImageDraw.Draw(im)
        W, H = im.size
        if i == 0 and settings.stamp:
            try:
                fnt = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 46)
            except Exception:
                fnt = ImageFont.load_default()
            x, y = int(W * 0.58), int(H * 0.06)
            tw = draw.textlength(settings.stamp, font=fnt)
            x = min(x, int(W - tw - 80))
            draw.rectangle([x - 18, y - 12, x + tw + 18, y + 64], outline=70, width=5)
            draw.text((x, y), settings.stamp, fill=70, font=fnt)
        if i == 0 and settings.handwriting:
            try:
                fnt = ImageFont.truetype("/System/Library/Fonts/Supplemental/Bradley Hand Bold.ttf", 52)
            except Exception:
                try:
                    fnt = ImageFont.truetype("/System/Library/Fonts/Supplemental/Georgia Italic.ttf", 48)
                except Exception:
                    fnt = ImageFont.load_default()
            draw.text((int(W * 0.12), int(H * 0.86)), settings.handwriting, fill=40, font=fnt)
        arr = np.array(im).astype(np.float32)
        nrng = np.random.RandomState(zlib.crc32(f"{seed}-{i}".encode()))
        arr += nrng.normal(0, 6 + 22 * settings.noise, arr.shape)
        speck = nrng.rand(*arr.shape) < 0.0008 * (0.3 + settings.noise)
        arr[speck] = 30
        arr = np.clip(arr * rnd.uniform(0.95, 1.0) + rnd.uniform(0, 8), 0, 255).astype(np.uint8)
        im = Image.fromarray(arr, "L")
        if settings.blur > 0:
            im = im.filter(ImageFilter.GaussianBlur(radius=settings.blur * 0.9))
        # recompress as JPEG like a scanner
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=72)
        out.append(Image.open(io.BytesIO(buf.getvalue())).convert("L"))
    out[0].save(pdf_out, "PDF", resolution=200.0, save_all=True, append_images=out[1:])


# ---------------------------------------------------------------- driver

def out_path(entry) -> Path:
    root = HELDBACK_ROOT if entry.root == "heldback" else ARCHIVE_ROOT
    return root / entry.path / entry.filename


def people_fn(world):
    def name(pid):
        p = world.people.get(pid)
        if not p:
            return pid
        # document properties usually show the Windows display name
        return p.canonical_name
    return name


def render_entry(entry, content, world, errors: list) -> Path | None:
    """Render one manifest entry from parsed content (pydantic) into the output tree."""
    from . import manifest as Mf
    style = Style(entry, world)
    people = people_fn(world)
    dst = out_path(entry)
    dst.parent.mkdir(parents=True, exist_ok=True)
    fam = Mf.family(entry)
    work = Path(tempfile.mkdtemp(prefix="mockgen_"))
    try:
        if fam == "deck":
            tmp = work / (Path(entry.filename).stem + ".pptx")
            render_deck(content, entry, style, people, tmp, errors)
        elif fam == "workbook":
            tmp = work / (Path(entry.filename).stem + ".xlsx")
            render_workbook(content, entry, style, people, tmp, errors)
        elif fam in ("document", "scan"):
            doc = content.document if fam == "scan" else content
            tmp = work / (Path(entry.filename).stem + ".docx")
            render_document(doc, entry, style, people, tmp, errors)
        else:
            errors.append(f"family {fam} is not rendered from content JSON")
            return None
        if errors:
            return None
        if entry.format in ("pptx", "xlsx", "docx"):
            shutil.copy(tmp, dst)
        elif entry.format in ("ppt", "xls", "doc"):
            conv = lo_convert(tmp, entry.format, work / "conv")
            shutil.copy(conv, dst)
        elif entry.format == "pdf":
            pdf = lo_convert(tmp, "pdf", work / "pdf")
            if fam == "scan":
                make_scan(pdf, dst, content.scan, entry.id)
                set_pdf_meta(dst, entry, people, "scan", "")
            else:
                shutil.copy(pdf, dst)
                set_pdf_meta(dst, entry, people, "native", getattr(content, "title_block", None) and content.title_block.title or entry.filename)
        set_mtime(dst, entry)
        return dst
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_export(entry, world, errors) -> Path | None:
    """PDF export of an Office file already rendered."""
    from . import manifest as Mf
    src_id = next(r.id for r in entry.related if r.rel == "export_of")
    src = Mf.entries_by_id().get(src_id)
    if not src:
        errors.append(f"export source {src_id} not in manifest")
        return None
    sp = out_path(src)
    if not sp.exists():
        errors.append(f"export source {src_id} not rendered yet")
        return None
    work = Path(tempfile.mkdtemp(prefix="mockgen_"))
    try:
        pdf = lo_convert(sp, "pdf", work)
        dst = out_path(entry)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(pdf, dst)
        title = ""
        try:
            c = json.loads((CONTENT / f"{src_id}.json").read_text())
            title = (c.get("properties") or {}).get("title") or (c.get("slides") or [{}])[0].get("title") or ""
        except Exception:
            pass
        set_pdf_meta(dst, entry, people_fn(world), "export", title)
        set_mtime(dst, entry)
        return dst
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_cmd(file_id=None, firm=None, all_=False) -> str:
    """Re-render accepted files from stored content JSON (no model calls)."""
    from . import manifest as Mf
    from .schemas import FAMILY_MODELS
    from .world import load_world
    from . import csvgen, state
    w = load_world([])
    es = Mf.load_entries([])
    if file_id:
        es = [e for e in es if e.id == file_id]
    elif firm:
        es = [e for e in es if e.unit == firm]
    elif not all_:
        return "give --file, --firm or --all"
    done, failed, skipped = 0, [], 0
    # sources before exports
    es = sorted(es, key=lambda e: (Mf.family(e) == "export", e.id))
    for e in es:
        fam = Mf.family(e)
        errs: list[str] = []
        if fam == "csv":
            p = csvgen.render_csv(e, w, errs)
        elif fam == "export":
            p = render_export(e, w, errs)
        else:
            cp = CONTENT / f"{e.id}.json"
            if not cp.exists():
                skipped += 1
                continue
            content = FAMILY_MODELS[fam].model_validate(json.loads(cp.read_text()))
            p = render_entry(e, content, w, errs)
        if errs or p is None:
            failed.append(f"{e.id}: {'; '.join(errs)[:300]}")
        else:
            done += 1
    return f"rendered {done}, skipped {skipped} (no content yet), failed {len(failed)}" + ("".join("\n  " + f for f in failed))
